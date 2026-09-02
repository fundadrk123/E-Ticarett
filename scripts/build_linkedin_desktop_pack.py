"""LinkedIn paylaşım paketini masaüstüne Word + görseller olarak hazırlar."""
from pathlib import Path
from shutil import copy2

from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT

DESKTOP = Path.home() / "Desktop"
PACK = DESKTOP / "LinkedIn-Paylasim-Paketi"
GORSeller = PACK / "Gorseller"
ROOT = Path(__file__).resolve().parent.parent
CAROUSEL = ROOT / "docs" / "linkedin" / "carousel"

POST_TEXT = """Yeni projem: Fullstack E-Ticaret Platformu 🚀

İnşaat malzemeleri ve hırdavat toptan satışına yönelik,
uçtan uca bir e-ticaret uygulaması geliştirdim.

Ne yaptım?
→ Ürün kataloğu, sepet, checkout (misafir + üye)
→ iyzico kredi kartı, havale, kapıda ödeme
→ JWT auth, kupon sistemi, sayısal stok yönetimi
→ Admin paneli: ürün, sipariş, kargo takibi, iade

Tech Stack:
Next.js 15 · React 19 · TypeScript · PostgreSQL 16 · Tailwind CSS · iyzico · JWT

37 REST API endpoint, 12 veritabanı tablosu, tek monolitik Next.js uygulaması.

💻 GitHub: https://github.com/fundadrk123/E-Ticarett

👉 Swipe — ekran görüntülerine bakın.

#nextjs #typescript #postgresql #fullstack #webdevelopment #eticaret #react #portfolio #iyzico"""

FIRST_COMMENT = """Teknik detay: Sepet hem localStorage hem PostgreSQL'de senkron çalışıyor.
Misafir kullanıcı localStorage'da, giriş yapan kullanıcıda sunucu sepeti ile birleştiriliyor.

Admin panelinde sipariş geldiğinde toast bildirimi + 20 sn'de bir polling var.

Kaynak kod GitHub'da açık — sorularınız varsa yorumlarda yanıtlarım 👇"""

IMAGE_ORDER = [
    ("01-cover.png", "1. Kapak — proje tanıtımı"),
    ("02-magaza.png", "2. Ana sayfa (tam ekran)"),
    ("03-katalog.png", "3. Ürün kataloğu"),
    ("04-urun-detay.png", "4. Ürün detay"),
    ("05-sepet.png", "5. Sepet"),
    ("06-odeme.png", "6. Ödeme / Checkout"),
    ("07-admin.png", "7. Admin giriş"),
    ("08-tech-stack.png", "8. Tech stack"),
    ("09-cta.png", "9. GitHub CTA"),
]

DESKTOP_NAMES = [
    "LinkedIn-1-KAPAK.png",
    "LinkedIn-2-ANA-SAYFA.png",
    "LinkedIn-3-KATALOG.png",
    "LinkedIn-4-URUN-DETAY.png",
    "LinkedIn-5-SEPET.png",
    "LinkedIn-6-ODEME.png",
    "LinkedIn-7-ADMIN.png",
    "LinkedIn-8-TECH-STACK.png",
    "LinkedIn-9-GITHUB.png",
]


def build_word():
    doc = Document()

    title = doc.add_heading("LinkedIn Paylaşım Paketi", 0)
    title.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER

    sub = doc.add_paragraph("Mertem Grup E-Ticaret · Fullstack Portföy")
    sub.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
    sub.runs[0].italic = True
    sub.runs[0].font.size = Pt(14)

    doc.add_paragraph()

    doc.add_heading("ADIM 1 — LinkedIn'de gönderi oluştur", level=1)
    steps = [
        "linkedin.com adresine gidin ve giriş yapın.",
        "Ana sayfada «Gönderi başlat» / «Start a post» alanına tıklayın.",
        "Aşağıdaki «PAYLAŞIM METNİ» bölümündeki metni kopyalayıp yapıştırın.",
        "«Görsel ekle» / «Add a photo» butonuna tıklayın.",
        "Bu klasördeki Gorseller klasöründen 01 ile 09 arası PNG dosyalarını SIRAYLA seçin.",
        "«Paylaş» / «Post» butonuna basın.",
        "Paylaşımdan hemen sonra «İLK YORUM» metnini kendi gönderinize yorum olarak yazın.",
    ]
    for i, s in enumerate(steps, 1):
        doc.add_paragraph(f"{i}. {s}", style="List Number")

    doc.add_heading("ADIM 2 — Görseller nerede?", level=1)
    p = doc.add_paragraph()
    p.add_run("Tüm görseller bu Word dosyasıyla aynı klasörde:\n").bold = False
    run = p.add_run(str(GORSeller))
    run.bold = True
    run.font.color.rgb = RGBColor(0x1D, 0x4E, 0xD8)

    doc.add_paragraph("Dosya sırası (LinkedIn'e bu sırayla yükleyin):", style="List Bullet")
    for fname, desc in IMAGE_ORDER:
        doc.add_paragraph(f"{fname} — {desc}", style="List Bullet")

    doc.add_page_break()

    doc.add_heading("PAYLAŞIM METNİ (kopyala-yapıştır)", level=1)
    box = doc.add_paragraph(POST_TEXT)
    for run in box.runs:
        run.font.size = Pt(11)

    doc.add_heading("İLK YORUM (paylaşımdan sonra yapıştır)", level=1)
    comment = doc.add_paragraph(FIRST_COMMENT)
    for run in comment.runs:
        run.font.size = Pt(11)

    doc.add_page_break()

    doc.add_heading("Carousel önizleme", level=1)
    doc.add_paragraph("LinkedIn'e yüklenecek slaytlar:")

    for fname, desc in IMAGE_ORDER:
        src = GORSeller / fname
        if src.exists():
            doc.add_heading(desc, level=2)
            try:
                doc.add_picture(str(src), width=Inches(4.5))
            except Exception:
                doc.add_paragraph(f"(Görsel: {fname})")

    out = PACK / "LinkedIn-Paylasim-Rehberi.docx"
    doc.save(str(out))
    return out


def main():
    PACK.mkdir(parents=True, exist_ok=True)
    GORSeller.mkdir(parents=True, exist_ok=True)

    copied = 0
    for (fname, _), desk_name in zip(IMAGE_ORDER, DESKTOP_NAMES):
        src = CAROUSEL / fname
        if src.exists():
            copy2(src, GORSeller / fname)
            copy2(src, DESKTOP / desk_name)
            copy2(src, PACK / desk_name)
            copied += 1

    # Ham tam ekran screenshotlar (1920x1080)
    raw = PACK / "Tam-Ekran-Screenshotlar"
    raw.mkdir(exist_ok=True)
    shot_src = ROOT / "docs" / "linkedin" / "screenshots"
    for f in sorted(shot_src.glob("0*.png")):
        copy2(f, raw / f.name)

    (PACK / "PAYLASIM-METNI.txt").write_text(
        POST_TEXT + "\n\n--- İLK YORUM ---\n\n" + FIRST_COMMENT,
        encoding="utf-8",
    )

    (PACK / "NASIL-KULLANILIR.txt").write_text(
        f"""LinkedIn Paylaşım Paketi
========================

1. LinkedIn-Paylasim-Rehberi.docx dosyasını açın (Word)
2. PAYLASIM-METNI.txt dosyasından metni kopyalayın VEYA Word içindeki metni kullanın
3. Gorseller klasöründeki 01-cover.png ... 09-cta.png dosyalarını LinkedIn'e sırayla yükleyin

Klasör yolu:
{PACK}

GitHub: https://github.com/fundadrk123/E-Ticarett
""",
        encoding="utf-8",
    )

    word_path = build_word()
    print(f"Klasor: {PACK}")
    print(f"Word:   {word_path}")
    print(f"Gorsel: {copied} adet kopyalandi")


if __name__ == "__main__":
    main()
